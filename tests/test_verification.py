"""Diff, verification, audit, and storage tests.

The load-bearing properties here are:

* diff on **content**, never bytes — proven by asserting that an unchanged
  workbook produces an empty diff despite differing bytes (ADR-0009)
* verification **can fail a run**, and a "saved" file is never "verified"
* formula checks are **static** and never claim recalculation (ADR-0011)
* secrets never reach the audit log
"""

from __future__ import annotations

import json
import zipfile
from pathlib import Path

import openpyxl
import pytest
from fixtures.workbooks import build, with_formula_damage

from app.audit import AuditLog, EventType, FileRunStore
from app.audit.redaction import Redactor as DirectRedactor
from app.contracts.base import UntrustedText, utc_now
from app.contracts.config import ExcelPilotConfig
from app.contracts.enums import Actor, VerificationStatus
from app.contracts.operations import Target
from app.contracts.pipeline import RunRecord
from app.contracts.verification import ReconciliationReport, ReconciliationResult
from app.diff import diff_paths
from app.storage import RunLocked
from app.verification import Verifier, describe
from app.verification import formulas as formula_checks
from app.workbook import file_sha256, opened, save_atomic

# --------------------------------------------------------------------------
# Diff
# --------------------------------------------------------------------------


class TestDiffIsContentBased:
    def test_unchanged_workbook_diff_is_empty_despite_differing_bytes(self, tmp_path: Path) -> None:
        """The central claim of ADR-0009, proven.

        A load/save cycle changes the file's bytes. If the diff compared bytes it
        would report changes here, and an operator would learn to ignore it.
        """
        source = build("monthly_sales", tmp_path / "s.xlsx", rows=15)
        workbook = openpyxl.load_workbook(source)
        resaved = tmp_path / "resaved.xlsx"
        save_atomic(workbook, resaved)
        workbook.close()

        assert file_sha256(source) != file_sha256(resaved), "bytes must differ"
        diff = diff_paths(source, resaved)
        assert diff.is_empty, f"content diff should be empty; got {diff.summary()}"

    def test_reports_value_change(self, tmp_path: Path) -> None:
        source = build("monthly_sales", tmp_path / "s.xlsx", rows=10)
        workbook = openpyxl.load_workbook(source)
        workbook["Sales"]["B2"] = "CHANGED"
        changed = tmp_path / "changed.xlsx"
        save_atomic(workbook, changed)
        workbook.close()

        diff = diff_paths(source, changed)
        assert diff.total_cell_changes == 1
        assert diff.cell_changes[0].coordinate == "B2"
        assert diff.cell_changes[0].change == "value"

    def test_reports_formula_replaced_by_literal(self, tmp_path: Path) -> None:
        source = build("monthly_sales", tmp_path / "s.xlsx", rows=10)
        workbook = openpyxl.load_workbook(source)
        workbook["Sales"]["G2"] = 999
        changed = tmp_path / "c.xlsx"
        save_atomic(workbook, changed)
        workbook.close()

        diff = diff_paths(source, changed)
        assert diff.total_formula_changes >= 1
        sheet = next(s for s in diff.sheet_diffs if s.name == "Sales")
        assert sheet.formulas_removed >= 1

    def test_reports_new_sheet_as_structural(self, tmp_path: Path) -> None:
        source = build("monthly_sales", tmp_path / "s.xlsx", rows=10)
        workbook = openpyxl.load_workbook(source)
        workbook.create_sheet("Added")
        changed = tmp_path / "c.xlsx"
        save_atomic(workbook, changed)
        workbook.close()

        diff = diff_paths(source, changed)
        assert diff.sheets_added == ["Added"]
        assert diff.structural_change is True

    def test_reports_removed_sheet(self, tmp_path: Path) -> None:
        source = build("monthly_sales", tmp_path / "s.xlsx", rows=10)
        workbook = openpyxl.load_workbook(source)
        del workbook["_Lookup"]
        changed = tmp_path / "c.xlsx"
        save_atomic(workbook, changed)
        workbook.close()

        diff = diff_paths(source, changed)
        assert "_Lookup" in diff.sheets_removed
        assert diff.structural_change is True

    def test_detects_rename_by_unchanged_content(self, tmp_path: Path) -> None:
        source = build("monthly_sales", tmp_path / "s.xlsx", rows=10)
        workbook = openpyxl.load_workbook(source)
        workbook["_Lookup"].title = "Reference"
        changed = tmp_path / "c.xlsx"
        save_atomic(workbook, changed)
        workbook.close()

        diff = diff_paths(source, changed)
        assert {"from": "_Lookup", "to": "Reference"} in diff.sheets_renamed
        assert diff.structural_change is True

    def test_ambiguous_rename_is_not_reported(self, tmp_path: Path) -> None:
        """Two identical sheets removed and re-added is ambiguous, so not a rename."""
        workbook = openpyxl.Workbook()
        first = workbook.active
        first.title = "A"
        first.append(["same"])
        second = workbook.create_sheet("B")
        second.append(["same"])
        source = tmp_path / "amb.xlsx"
        workbook.save(source)
        workbook.close()

        workbook = openpyxl.load_workbook(source)
        del workbook["A"]
        del workbook["B"]
        workbook.create_sheet("C").append(["same"])
        workbook.create_sheet("D").append(["same"])
        changed = tmp_path / "amb2.xlsx"
        save_atomic(workbook, changed)
        workbook.close()

        diff = diff_paths(source, changed)
        assert diff.sheets_renamed == []

    def test_reports_defined_name_change(self, tmp_path: Path) -> None:
        source = build("table", tmp_path / "t.xlsx")
        workbook = openpyxl.load_workbook(source)
        from openpyxl.workbook.defined_name import DefinedName

        workbook.defined_names.add(DefinedName("BrandNew", attr_text="Orders!$A$2:$A$5"))
        changed = tmp_path / "c.xlsx"
        save_atomic(workbook, changed)
        workbook.close()

        diff = diff_paths(source, changed)
        assert "BrandNew" in diff.defined_names_added
        assert diff.structural_change is True

    def test_change_list_is_capped_but_totals_stay_exact(self, tmp_path: Path) -> None:
        from app.diff.engine import MAX_CELL_CHANGES

        total = MAX_CELL_CHANGES + 200
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Big"
        for index in range(total):
            # +1 so every value differs after doubling; 0 * 2 == 0 would not.
            sheet.cell(row=index + 1, column=1, value=index + 1)
        source = tmp_path / "big.xlsx"
        workbook.save(source)
        workbook.close()

        workbook = openpyxl.load_workbook(source)
        target = workbook["Big"]
        for index in range(total):
            target.cell(row=index + 1, column=1, value=(index + 1) * 2)
        changed = tmp_path / "big2.xlsx"
        save_atomic(workbook, changed)
        workbook.close()

        diff = diff_paths(source, changed)
        assert len(diff.cell_changes) <= MAX_CELL_CHANGES
        assert diff.cell_changes_truncated is True
        # The total is exact even though the list is capped.
        assert diff.total_cell_changes == total

    def test_diff_records_both_hashes(self, tmp_path: Path) -> None:
        source = build("minimal", tmp_path / "m.xlsx")
        copy = tmp_path / "copy.xlsx"
        copy.write_bytes(source.read_bytes())
        diff = diff_paths(source, copy)
        assert diff.before_hash == file_sha256(source)
        assert diff.after_hash == file_sha256(copy)

    def test_notes_disclose_the_limitation(self, tmp_path: Path) -> None:
        source = build("minimal", tmp_path / "m.xlsx")
        diff = diff_paths(source, source)
        assert any("not file bytes" in note for note in diff.notes)
        assert any("style index" in note for note in diff.notes)


# --------------------------------------------------------------------------
# Formula verification
# --------------------------------------------------------------------------


class TestFormulaChecks:
    def test_collects_formulas(self, tmp_path: Path) -> None:
        path = build("monthly_sales", tmp_path / "s.xlsx", rows=5)
        with opened(path) as workbook:
            formulas = formula_checks.collect_formulas(workbook)
        assert len(formulas) >= 5
        assert all(formula.startswith("=") for _, _, formula in formulas)

    def test_detects_broken_reference(self, tmp_path: Path) -> None:
        workbook = with_formula_damage(rows=10)
        path = tmp_path / "d.xlsx"
        workbook.save(path)
        workbook.close()
        with opened(path) as loaded:
            result = formula_checks.check_broken_references(loaded)
        assert result.failed is True
        assert "#REF!" in result.message or "broken" in result.message

    def test_clean_workbook_passes(self, tmp_path: Path) -> None:
        path = build("monthly_sales", tmp_path / "s.xlsx", rows=5)
        with opened(path) as workbook:
            result = formula_checks.check_broken_references(workbook)
        assert result.status is VerificationStatus.PASSED

    def test_detects_external_reference_as_warning(self, tmp_path: Path) -> None:
        workbook = with_formula_damage(rows=10)
        path = tmp_path / "d.xlsx"
        workbook.save(path)
        workbook.close()
        with opened(path) as loaded:
            result = formula_checks.check_external_references(loaded)
        assert result.status is VerificationStatus.WARNING

    def test_detects_inconsistent_column_pattern(self, tmp_path: Path) -> None:
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "D"
        sheet.append(["A", "B"])
        for index in range(10):
            sheet.append([index, None])
            sheet.cell(row=index + 2, column=2, value=f"=A{index + 2}*2")
        # One formula with a different structure.
        sheet.cell(row=3, column=2, value="=SUM(A1:A10)")
        path = tmp_path / "inc.xlsx"
        workbook.save(path)
        workbook.close()
        with opened(path) as loaded:
            result = formula_checks.check_column_consistency(loaded)
        assert result.status is VerificationStatus.WARNING

    def test_consistent_column_passes(self, tmp_path: Path) -> None:
        path = build("monthly_sales", tmp_path / "s.xlsx", rows=20)
        with opened(path) as workbook:
            result = formula_checks.check_column_consistency(workbook)
        assert result.status is VerificationStatus.PASSED

    def test_detects_formula_loss(self, tmp_path: Path) -> None:
        path = build("monthly_sales", tmp_path / "s.xlsx", rows=10)
        with opened(path) as loaded:
            before = {
                (sheet, coordinate): formula
                for sheet, coordinate, formula in formula_checks.collect_formulas(loaded)
            }
        workbook = openpyxl.load_workbook(path)
        for coordinate in ("G2", "G3", "G4"):
            workbook["Sales"][coordinate] = 1
        changed = tmp_path / "c.xlsx"
        save_atomic(workbook, changed)
        workbook.close()
        with opened(changed) as loaded:
            result, lost = formula_checks.detect_formula_loss(before, loaded)
        assert result.failed is True
        assert len(lost) == 3

    def test_detects_hardcoded_replacement(self, tmp_path: Path) -> None:
        path = build("monthly_sales", tmp_path / "s.xlsx", rows=10)
        with opened(path) as loaded:
            before = {
                (sheet, coordinate): formula
                for sheet, coordinate, formula in formula_checks.collect_formulas(loaded)
            }
        workbook = openpyxl.load_workbook(path)
        workbook["Sales"]["G2"] = 12345
        changed = tmp_path / "c.xlsx"
        save_atomic(workbook, changed)
        workbook.close()
        with opened(changed) as loaded:
            result = formula_checks.detect_hardcoded_replacements(before, loaded)
        assert result.failed is True
        assert "12345" in result.details["issues"][0]

    def test_reports_unresolvable_dynamic_formulas(self, tmp_path: Path) -> None:
        path = build("formulas_only", tmp_path / "f.xlsx", rows=5)
        with opened(path) as workbook:
            unresolvable = formula_checks.unresolvable_formulas(workbook)
        # The fixture's formulas are plain arithmetic, so nothing is unresolvable.
        assert unresolvable == []

    def test_indirect_is_reported_as_unresolvable(self, tmp_path: Path) -> None:
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "S"
        sheet.append(["A"])
        sheet.append([1])
        sheet["B1"] = '=INDIRECT("A1")'
        path = tmp_path / "i.xlsx"
        workbook.save(path)
        workbook.close()
        with opened(path) as loaded:
            unresolvable = formula_checks.unresolvable_formulas(loaded)
        assert len(unresolvable) == 1
        assert "INDIRECT" in unresolvable[0].detail


# --------------------------------------------------------------------------
# Verification
# --------------------------------------------------------------------------


class TestVerifier:
    def test_passing_verification_of_a_no_op(self, tmp_path: Path) -> None:
        source = build("monthly_sales", tmp_path / "s.xlsx", rows=15)
        workbook = openpyxl.load_workbook(source)
        output = tmp_path / "out.xlsx"
        save_atomic(workbook, output)
        workbook.close()

        result = Verifier().verify("run-1", output, before_path=source)
        assert result.passed is True
        assert result.recalculated is False
        assert result.static_formula_checks is True

    def test_verification_fails_when_a_sheet_disappears(self, tmp_path: Path) -> None:
        source = build("monthly_sales", tmp_path / "s.xlsx", rows=10)
        workbook = openpyxl.load_workbook(source)
        del workbook["_Lookup"]
        output = tmp_path / "out.xlsx"
        save_atomic(workbook, output)
        workbook.close()

        result = Verifier().verify("run-1", output, before_path=source)
        assert result.passed is False
        removed = [c for c in result.structural if c.name == "structure_no_sheet_removed"]
        assert removed[0].failed is True

    def test_verification_fails_when_formulas_are_destroyed(self, tmp_path: Path) -> None:
        source = build("monthly_sales", tmp_path / "s.xlsx", rows=10)
        workbook = openpyxl.load_workbook(source)
        for row in range(2, 12):
            workbook["Sales"].cell(row=row, column=7, value=1)
        output = tmp_path / "out.xlsx"
        save_atomic(workbook, output)
        workbook.close()

        result = Verifier().verify("run-1", output, before_path=source)
        assert result.passed is False
        preservation = [c for c in result.formula if c.name == "formula_preservation"]
        assert preservation[0].failed is True

    def test_missing_output_fails(self, tmp_path: Path) -> None:
        result = Verifier().verify("run-1", tmp_path / "does-not-exist.xlsx")
        assert result.passed is False
        assert any("does not exist" in note for note in result.notes)

    def test_corrupt_output_fails_the_file_check(self, tmp_path: Path) -> None:
        source = build("minimal", tmp_path / "m.xlsx")
        output = tmp_path / "out.xlsx"
        with zipfile.ZipFile(source) as archive, zipfile.ZipFile(output, "w") as target:
            for item in archive.infolist():
                target.writestr(item, archive.read(item.filename))
        # Truncate to make it unreadable.
        output.write_bytes(b"not a workbook")

        result = Verifier().verify("run-1", output, before_path=source)
        assert result.passed is False

    def test_reconciliation_failure_fails_the_run(self, tmp_path: Path) -> None:
        source = build("monthly_sales", tmp_path / "s.xlsx", rows=10)
        report = ReconciliationReport(
            run_id="run-1",
            results=[
                ReconciliationResult(
                    name="total",
                    status="failed",
                    expected=100,
                    actual=50,
                    variance=50,
                    explanation="does not match",
                )
            ],
            failed_count=1,
        )
        result = Verifier().verify("run-1", source, reconciliation=report)
        assert result.passed is False

    def test_never_claims_recalculation(self, tmp_path: Path) -> None:
        source = build("formulas_only", tmp_path / "f.xlsx", rows=8)
        result = Verifier().verify("run-1", source, before_path=source)
        assert result.recalculated is False
        assert result.reconciliation is None or result.reconciliation.recalculated is False
        text = describe(result)
        assert "recalculated:      False" in text

    def test_a_saved_file_is_not_a_verified_outcome(self, tmp_path: Path) -> None:
        """The distinction the whole verification layer exists to preserve."""
        source = build("monthly_sales", tmp_path / "s.xlsx", rows=10)
        workbook = openpyxl.load_workbook(source)
        del workbook["_Lookup"]
        output = tmp_path / "out.xlsx"
        save_atomic(workbook, output)
        workbook.close()

        assert output.exists(), "the file was written"
        result = Verifier().verify("run-1", output, before_path=source)
        assert result.passed is False, "but the outcome was not verified"
        assert "do not treat the output as correct" in describe(result)

    def test_data_checks_report_row_loss(self, tmp_path: Path) -> None:
        source = build("monthly_sales", tmp_path / "s.xlsx", rows=30)
        workbook = openpyxl.load_workbook(source)
        worksheet = workbook["Sales"]
        worksheet.delete_rows(2, 25)
        output = tmp_path / "out.xlsx"
        save_atomic(workbook, output)
        workbook.close()

        result = Verifier().verify(
            "run-1",
            output,
            before_path=source,
            data_targets=[Target(sheet="Sales")],
        )
        row_check = [c for c in result.data if c.name == "data_row_count"]
        assert row_check and row_check[0].status in {
            VerificationStatus.FAILED,
            VerificationStatus.WARNING,
        }

    def test_anomalies_detect_sheet_removal(self, tmp_path: Path) -> None:
        source = build("monthly_sales", tmp_path / "s.xlsx", rows=10)
        workbook = openpyxl.load_workbook(source)
        del workbook["_Lookup"]
        output = tmp_path / "out.xlsx"
        save_atomic(workbook, output)
        workbook.close()

        result = Verifier().verify("run-1", output, before_path=source)
        kinds = {anomaly.kind.value for anomaly in result.anomalies}
        assert "structural_change" in kinds

    def test_planned_vs_actual_divergence_is_flagged(self, tmp_path: Path) -> None:
        source = build("monthly_sales", tmp_path / "s.xlsx", rows=10)
        result = Verifier().verify(
            "run-1",
            source,
            before_path=source,
            planned_cells_changed=100,
            actual_cells_changed=900,
        )
        kinds = {anomaly.kind.value for anomaly in result.anomalies}
        assert "planned_vs_actual_divergence" in kinds

    def test_expected_sheets_checked(self, tmp_path: Path) -> None:
        source = build("monthly_sales", tmp_path / "s.xlsx", rows=10)
        result = Verifier().verify(
            "run-1", source, before_path=source, expected_sheets=["Sales", "Nonexistent"]
        )
        assert result.passed is False
        check = [c for c in result.structural if c.name == "structure_expected_sheets"][0]
        assert check.failed is True

    def test_describe_includes_every_family(self, tmp_path: Path) -> None:
        source = build("monthly_sales", tmp_path / "s.xlsx", rows=10)
        result = Verifier().verify(
            "run-1",
            source,
            before_path=source,
            data_targets=[Target(sheet="Sales")],
            reconciliation=ReconciliationReport(run_id="run-1", results=[], passed_count=0),
        )
        text = describe(result)
        assert "Structural:" in text
        assert "Formula:" in text
        assert "Reconciliation:" in text


# --------------------------------------------------------------------------
# Audit
# --------------------------------------------------------------------------


class TestRedaction:
    def test_registered_secret_is_masked(self, monkeypatch: pytest.MonkeyPatch) -> None:
        monkeypatch.setenv("TYPESAFE_API_KEY", "tsk-supersecretvalue12345")
        redactor = DirectRedactor()
        masked = redactor.mask("calling with tsk-supersecretvalue12345 now")
        assert "supersecret" not in masked
        assert "[REDACTED" in masked

    @pytest.mark.parametrize(
        ("value", "kind"),
        [
            ("sk-abcdefghijklmnopqrstuvwxyz012345", "openai_key"),
            ("sk-ant-abcdefghijklmnopqrstuvwxyz", "anthropic_key"),
            ("AKIAIOSFODNN7EXAMPLE", "aws_access_key"),
            ("AIzaSyAbcdefghijklmnopqrstuvwxyz0123456", "google_api_key"),
        ],
    )
    def test_shaped_secrets_are_masked_without_being_registered(
        self, value: str, kind: str, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """A secret ExcelPilot was never told about must still not leak."""
        for name in list(__import__("os").environ):
            if name.endswith("_API_KEY") or "TOKEN" in name or "SECRET" in name:
                monkeypatch.delenv(name, raising=False)
        masked = DirectRedactor().mask(f"value is {value} end")
        assert value not in masked
        assert kind in masked

    def test_sensitive_keys_masked_wholesale(self, monkeypatch: pytest.MonkeyPatch) -> None:
        redactor = DirectRedactor()
        result = redactor.redact(
            {"api_key": "anything", "password": "hunter2", "Authorization": "Bearer x", "ok": 1}
        )
        assert result["api_key"] == "[REDACTED:sensitive_key]"
        assert result["password"] == "[REDACTED:sensitive_key]"
        assert result["Authorization"] == "[REDACTED:sensitive_key]"
        assert result["ok"] == 1

    def test_nested_structures_are_redacted(self, monkeypatch: pytest.MonkeyPatch) -> None:
        monkeypatch.setenv("OPENAI_API_KEY", "sk-abcdefghijklmnopqrstuvwxyz012345")
        result = DirectRedactor().redact(
            {"outer": {"inner": [{"token": "x", "note": "sk-abcdefghijklmnopqrstuvwxyz012345"}]}}
        )
        serialised = json.dumps(result)
        assert "sk-abcdefghijklmnopqrstuvwxyz012345" not in serialised

    def test_private_key_block_is_masked(self, monkeypatch: pytest.MonkeyPatch) -> None:
        block = "-----BEGIN RSA PRIVATE KEY-----\nMIIabc\n-----END RSA PRIVATE KEY-----"
        assert "MIIabc" not in DirectRedactor().mask(block)

    def test_connection_string_password_is_masked(self, monkeypatch: pytest.MonkeyPatch) -> None:
        masked = DirectRedactor().mask("postgres://user:hunter2@localhost/db")
        assert "hunter2" not in masked

    def test_redactor_never_raises(self, monkeypatch: pytest.MonkeyPatch) -> None:
        class Unprintable:
            def __str__(self) -> str:
                raise RuntimeError("boom")

        assert DirectRedactor().redact({"x": Unprintable()}) is not None

    def test_depth_is_capped(self, monkeypatch: pytest.MonkeyPatch) -> None:
        deep: dict = {"a": 1}
        for _ in range(30):
            deep = {"n": deep}
        result = DirectRedactor().redact(deep)
        assert "too_deep" in json.dumps(result)


class TestAuditLog:
    def test_appends_events(self, tmp_path: Path) -> None:
        with AuditLog(tmp_path / "audit.jsonl") as log:
            log.emit(Actor.SYSTEM, EventType.RUN_CREATED, {"run_id": "r1"})
            log.emit(Actor.POLICY, EventType.POLICY_EVALUATED, {"run_id": "r1", "outcome": "allow"})
        events = list(AuditLog(tmp_path / "audit.jsonl").read())
        assert len(events) == 2
        assert events[0].seq == 1
        assert events[1].seq == 2

    def test_is_append_only_across_instances(self, tmp_path: Path) -> None:
        path = tmp_path / "audit.jsonl"
        with AuditLog(path) as log:
            log.emit(Actor.SYSTEM, EventType.RUN_CREATED, {"run_id": "r1"})
        with AuditLog(path) as log:
            log.emit(Actor.SYSTEM, EventType.RUN_COMPLETED, {"run_id": "r1"})
        events = list(AuditLog(path).read())
        assert [e.seq for e in events] == [1, 2]
        assert events[0].event_type == EventType.RUN_CREATED, "history must be preserved"

    def test_corrupt_trailing_line_is_skipped(self, tmp_path: Path) -> None:
        path = tmp_path / "audit.jsonl"
        with AuditLog(path) as log:
            log.emit(Actor.SYSTEM, EventType.RUN_CREATED, {"run_id": "r1"})
        with path.open("a", encoding="utf-8") as handle:
            handle.write('{"seq": 2, "trunc')  # simulate a crash mid-write
        events = list(AuditLog(path).read())
        assert len(events) == 1, "a partial line must not make the log unreadable"

    def test_secrets_never_reach_disk(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        monkeypatch.setenv("TYPESAFE_API_KEY", "tsk-leakcanary-987654321")
        path = tmp_path / "audit.jsonl"
        with AuditLog(path) as log:
            log.emit(
                Actor.JEV,
                EventType.JEV_DECIDED,
                {
                    "run_id": "r1",
                    "api_key": "tsk-leakcanary-987654321",
                    "note": "used tsk-leakcanary-987654321 to call",
                },
            )
        raw = path.read_text()
        assert "leakcanary" not in raw, "a secret reached the audit log"
        assert "[REDACTED" in raw

    def test_summary_counts_by_type_and_actor(self, tmp_path: Path) -> None:
        with AuditLog(tmp_path / "a.jsonl") as log:
            log.emit(Actor.SYSTEM, EventType.RUN_CREATED, {"run_id": "r"})
            log.emit(Actor.SYSTEM, EventType.RUN_COMPLETED, {"run_id": "r"})
            log.emit(Actor.POLICY, EventType.POLICY_EVALUATED, {"run_id": "r"})
        summary = AuditLog(tmp_path / "a.jsonl").summary()
        assert summary["events"] == 3
        assert summary["by_type"][EventType.RUN_CREATED] == 1
        assert summary["by_actor"]["system"] == 2


# --------------------------------------------------------------------------
# Storage
# --------------------------------------------------------------------------


class TestFileRunStore:
    @pytest.fixture
    def store(self, tmp_path: Path) -> FileRunStore:
        return FileRunStore(ExcelPilotConfig(workspace_root=str(tmp_path)))

    def test_creates_run_directory(self, store: FileRunStore) -> None:
        paths = store.create("run-abc")
        assert paths.root.exists()
        assert paths.output_dir.exists()

    def test_duplicate_create_fails(self, store: FileRunStore) -> None:
        store.create("run-abc")
        with pytest.raises(Exception, match="already exists"):
            store.create("run-abc")

    def test_rejects_path_traversal_in_run_id(self, store: FileRunStore) -> None:
        for bad in ("../escape", "a/b", "..", ".hidden"):
            with pytest.raises(Exception, match="invalid run id"):
                store.paths(bad)

    def test_snapshot_is_byte_identical(self, store: FileRunStore, tmp_path: Path) -> None:
        source = build("monthly_sales", tmp_path / "s.xlsx", rows=10)
        store.create("run-1")
        snapshot = store.snapshot("run-1", source)
        assert file_sha256(snapshot) == file_sha256(source)

    def test_run_record_roundtrip(self, store: FileRunStore) -> None:
        record = RunRecord(
            run_id="run-1",
            created_at=utc_now().isoformat(),
            raw_task=UntrustedText("clean the data", provenance="user_task"),
            source_path="/tmp/a.xlsx",
            source_name="a.xlsx",
            source_hash="abc",
        )
        store.write_run(record)
        loaded = store.read_run("run-1")
        assert loaded.run_id == "run-1"
        assert loaded.raw_task.text == "clean the data"

    def test_missing_run_lists_alternatives(self, store: FileRunStore) -> None:
        store.write_run(
            RunRecord(
                run_id="run-existing",
                created_at=utc_now().isoformat(),
                raw_task=UntrustedText("x", provenance="user_task"),
                source_path="/tmp/a.xlsx",
                source_name="a.xlsx",
                source_hash="h",
            )
        )
        with pytest.raises(Exception, match="run-existing"):
            store.read_run("run-nope")

    def test_corrupt_run_record_is_reported(self, store: FileRunStore) -> None:
        store.create("run-1")
        store.paths("run-1").run_file.write_text("{not json")
        with pytest.raises(Exception, match="corrupt"):
            store.read_run("run-1")

    def test_list_and_recent(self, store: FileRunStore) -> None:
        for index in range(3):
            store.write_run(
                RunRecord(
                    run_id=f"run-{index}",
                    created_at=utc_now().isoformat(),
                    raw_task=UntrustedText("x", provenance="user_task"),
                    source_path="/tmp/a.xlsx",
                    source_name="a.xlsx",
                    source_hash="h",
                )
            )
        assert len(store.list_runs()) == 3
        assert len(store.recent(limit=2)) == 2

    def test_lock_is_exclusive(self, store: FileRunStore) -> None:
        with store.lock("run-1"), pytest.raises(RunLocked), store.lock("run-1"):
            pass
        # Released afterwards.
        with store.lock("run-1"):
            pass

    def test_update_stamps_completion(self, store: FileRunStore) -> None:
        from app.contracts.enums import RunOutcome

        record = RunRecord(
            run_id="run-1",
            created_at=utc_now().isoformat(),
            raw_task=UntrustedText("x", provenance="user_task"),
            source_path="/tmp/a.xlsx",
            source_name="a.xlsx",
            source_hash="h",
        )
        updated = store.update_run(record.model_copy(update={"outcome": RunOutcome.SUCCEEDED}))
        assert updated.completed_at is not None

    def test_prune_keeps_newest(self, store: FileRunStore) -> None:
        for index in range(5):
            store.write_run(
                RunRecord(
                    run_id=f"run-{index}",
                    created_at=utc_now().isoformat(),
                    raw_task=UntrustedText("x", provenance="user_task"),
                    source_path="/tmp/a.xlsx",
                    source_name="a.xlsx",
                    source_hash="h",
                )
            )
        removed = store.prune(keep=2)
        assert len(removed) == 3
        assert len(store.list_runs()) == 2

    def test_output_path_is_inside_the_run(self, store: FileRunStore) -> None:
        path = store.output_path("run-1", "out.xlsx")
        assert path.parent.name == "output"
        assert "run-1" in str(path)

    def test_record_output_hashes_the_artefact(self, store: FileRunStore, tmp_path: Path) -> None:
        artefact = tmp_path / "a.xlsx"
        artefact.write_bytes(b"content")
        record = store.record_output(artefact)
        assert record["hash"] == file_sha256(artefact)
        assert record["size_bytes"] == 7
