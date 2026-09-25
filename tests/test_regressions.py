"""Regression tests for bugs found by running the pipeline end to end.

Each of these was a real defect discovered during implementation, not a
theoretical one. They are collected here so the specific failure cannot return,
with a comment explaining what went wrong and why the test asserts what it does.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest
from fixtures.workbooks import build

from app.app import RunOrchestrator
from app.contracts.base import UntrustedText
from app.contracts.config import ExcelPilotConfig
from app.contracts.enums import InterpretationVerdict, RunOutcome
from app.decisions import MockJevAdapter
from app.planner import DeterministicPlanner
from app.verification import formulas as formula_checks
from app.workbook import opened, save_atomic

#: The compound request that exposed the dedupe-key bug.
COMPOUND = (
    "normalise the Customer column, remove duplicate invoices, and create a summary by Region"
)


@pytest.fixture
def sales(tmp_path: Path) -> Path:
    return build("monthly_sales", tmp_path / "monthly_sales.xlsx", rows=40)


@pytest.fixture
def orchestrator(tmp_path: Path) -> RunOrchestrator:
    return RunOrchestrator(
        ExcelPilotConfig(workspace_root=str(tmp_path)), jev=MockJevAdapter("approve")
    )


class TestDedupeKeyIsNotGuessed:
    """A wrong dedupe key silently deletes rows that were not duplicates.

    The first version of the planner took columns from anywhere in the request, so
    the compound request below produced the key ``[Region, Customer]`` — a valid
    pair of real columns, which therefore passed *every* validation, while
    collapsing 38 of 42 rows. Only the row-count anomaly and verification caught
    it. The correct fix is not to emit the key at all unless the request states
    it, so that is what is asserted here.
    """

    def test_compound_request_uses_the_named_key(self, sales: Path) -> None:
        from app.workbook import inspect_workbook

        plan = DeterministicPlanner().plan(
            UntrustedText(COMPOUND, provenance="user_task"), inspect_workbook(sales)
        )
        dedupe = next(op for op in plan.operations if op.operation.value == "remove_duplicates")
        assert dedupe.keys == ["InvoiceId"], f"expected the invoice key, got {dedupe.keys}"

    def test_dedupe_does_not_borrow_another_intents_columns(self, sales: Path) -> None:
        """Normalise's 'Customer' and summary's 'Region' must not leak into dedupe."""
        from app.workbook import inspect_workbook

        plan = DeterministicPlanner().plan(
            UntrustedText(COMPOUND, provenance="user_task"), inspect_workbook(sales)
        )
        dedupe = next(op for op in plan.operations if op.operation.value == "remove_duplicates")
        assert "Customer" not in dedupe.keys
        assert "Region" not in dedupe.keys

    def test_explicit_key_is_honoured(self, sales: Path) -> None:
        from app.workbook import inspect_workbook

        plan = DeterministicPlanner().plan(
            UntrustedText("remove duplicate records by Customer", provenance="user_task"),
            inspect_workbook(sales),
        )
        dedupe = next(op for op in plan.operations if op.operation.value == "remove_duplicates")
        assert dedupe.keys == ["Customer"]

    def test_unstated_key_is_refused_not_guessed(self, tmp_path: Path) -> None:
        """With no key stated anywhere, the planner must refuse rather than invent one."""
        from app.workbook import inspect_workbook

        # A workbook with no column resembling an identifier or a concept word.
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "S"
        sheet.append(["Alpha", "Beta"])
        for index in range(5):
            sheet.append([f"a{index}", f"b{index}"])
        path = tmp_path / "opaque.xlsx"
        workbook.save(path)
        workbook.close()

        plan = DeterministicPlanner().plan(
            UntrustedText("remove duplicate records", provenance="user_task"),
            inspect_workbook(path),
        )
        assert plan.understanding.interpretation is InterpretationVerdict.REQUIRES_USER_INPUT
        assert any("will not guess" in item for item in plan.understanding.missing_information), (
            plan.understanding.missing_information
        )

    def test_real_dedupe_removes_only_the_actual_duplicates(
        self, sales: Path, orchestrator: RunOrchestrator
    ) -> None:
        """The end-to-end consequence: 2 duplicates, not 38."""
        before = openpyxl.load_workbook(sales)
        rows_before = before["Sales"].max_row
        before.close()

        result = orchestrator.run(
            sales, UntrustedText(COMPOUND, provenance="user_task"), approve=True
        )
        assert result.outcome is RunOutcome.SUCCEEDED, result.error

        with opened(str(result.output_path)) as after:
            rows_after = after["Sales"].max_row
        # The fixture contains 40 rows plus 2 exact duplicates.
        assert rows_before - rows_after == 2, (
            f"expected exactly the 2 duplicate rows to be removed, "
            f"but {rows_before} -> {rows_after}"
        )

    def test_customer_values_are_normalised(
        self, sales: Path, orchestrator: RunOrchestrator
    ) -> None:
        result = orchestrator.run(
            sales, UntrustedText(COMPOUND, provenance="user_task"), approve=True
        )
        assert result.succeeded, result.error
        with opened(str(result.output_path)) as after:
            customers = [after["Sales"].cell(row=row, column=2).value for row in range(2, 12)]
        assert all(value == value.strip() for value in customers if value), customers


class TestSheetNameIsNotConfusedWithAnOperationWord:
    """A sheet named 'Summary' must not capture 'create a summary by Region'.

    The planner originally matched sheet names by substring, so the word
    "summary" in the request selected the small `Summary` sheet — which has no
    Region column — and the plan was refused. Sheet references now require a
    capitalised whole word.
    """

    def test_operation_word_does_not_select_the_summary_sheet(self, sales: Path) -> None:
        from app.workbook import inspect_workbook

        plan = DeterministicPlanner().plan(
            UntrustedText("create a summary by Region", provenance="user_task"),
            inspect_workbook(sales),
        )
        summary = next(op for op in plan.operations if op.operation.value == "create_summary")
        assert summary.target.sheet == "Sales"

    def test_explicitly_named_summary_sheet_is_still_reachable(self, sales: Path) -> None:
        from app.workbook import inspect_workbook

        plan = DeterministicPlanner().plan(
            UntrustedText("read the Summary sheet", provenance="user_task"),
            inspect_workbook(sales),
        )
        read = next(op for op in plan.operations if op.operation.value == "read_range")
        assert read.target.sheet == "Summary"


class TestColumnMentionIsLocalToItsClause:
    """A column phrase must be captured as a single token.

    An earlier regex allowed spaces, so in "normalise the Customer column, ..."
    it captured "normalise the Customer" — the engine takes the earliest viable
    start — and the real column name was never found.
    """

    @pytest.mark.parametrize(
        ("phrase", "expected"),
        [
            ("normalise the Customer column", "Customer"),
            ("clean the Notes field", "Notes"),
        ],
    )
    def test_mentions_resolve(self, sales: Path, phrase: str, expected: str) -> None:
        from app.workbook import inspect_workbook

        plan = DeterministicPlanner().plan(
            UntrustedText(phrase, provenance="user_task"), inspect_workbook(sales)
        )
        targets = [op for op in plan.operations if hasattr(op, "columns")]
        assert any(expected in op.columns for op in targets), (
            f"{phrase!r} did not resolve {expected!r}: {[op.columns for op in targets]}"
        )

    def test_sort_column_resolves(self, sales: Path) -> None:
        """A sort names its column in ``by_columns``, not ``columns``."""
        from app.workbook import inspect_workbook

        plan = DeterministicPlanner().plan(
            UntrustedText("sort by Amount", provenance="user_task"), inspect_workbook(sales)
        )
        sort = next(op for op in plan.operations if op.operation.value == "sort_range")
        assert sort.by_columns == ["Amount"]


class TestSheetNameDoesNotResolveToAColumn:
    """A sheet can never name a column on itself.

    "normalise the Customer column on the Sales sheet" leaked the token "Sales"
    into column resolution, where the synonym table mapped it to the *Amount*
    column — "sales" is listed as a synonym for amount. The request silently
    normalised two columns instead of one.
    """

    def test_sheet_name_is_not_treated_as_a_column(self, sales: Path) -> None:
        from app.workbook import inspect_workbook

        plan = DeterministicPlanner().plan(
            UntrustedText(
                "normalise the Customer column on the Sales sheet", provenance="user_task"
            ),
            inspect_workbook(sales),
        )
        normalize = next(op for op in plan.operations if op.operation.value == "normalize_values")
        assert normalize.columns == ["Customer"], (
            f"only the named column should be touched, got {normalize.columns}"
        )
        assert "Amount" not in normalize.columns

    def test_a_sheet_named_after_a_synonym_still_works(self, tmp_path: Path) -> None:
        """The same request against a sheet genuinely named 'Total'."""
        from app.workbook import inspect_workbook

        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Total"
        sheet.append(["Customer", "Region"])
        for index in range(5):
            sheet.append([f"c{index}", f"r{index}"])
        path = tmp_path / "total.xlsx"
        workbook.save(path)
        workbook.close()

        plan = DeterministicPlanner().plan(
            UntrustedText(
                "normalise the Customer column on the Total sheet", provenance="user_task"
            ),
            inspect_workbook(path),
        )
        normalize = next(op for op in plan.operations if op.operation.value == "normalize_values")
        assert normalize.target.sheet == "Total"
        assert normalize.columns == ["Customer"]


class TestMutatingRequestsMustBeSpecific:
    """A request that names neither a sheet nor a column is refused.

    "Tidy it up a bit" used to produce a plan that normalised every column of
    whichever sheet happened to be largest — a change nobody asked for.
    """

    def test_vague_mutating_request_is_refused(self, sales: Path) -> None:
        from app.workbook import inspect_workbook

        plan = DeterministicPlanner().plan(
            UntrustedText("tidy it up a bit", provenance="user_task"), inspect_workbook(sales)
        )
        assert plan.understanding.interpretation is InterpretationVerdict.REQUIRES_USER_INPUT
        assert any("will not guess" in item for item in plan.understanding.missing_information)

    def test_a_named_column_is_enough(self, sales: Path) -> None:
        from app.workbook import inspect_workbook

        plan = DeterministicPlanner().plan(
            UntrustedText("tidy the Customer names", provenance="user_task"),
            inspect_workbook(sales),
        )
        assert plan.understanding.interpretation is InterpretationVerdict.SUFFICIENTLY_CLEAR

    def test_a_concept_named_key_is_enough(self, sales: Path) -> None:
        """ "remove duplicate invoices" names no sheet but resolves its key."""
        from app.workbook import inspect_workbook

        plan = DeterministicPlanner().plan(
            UntrustedText("remove duplicate invoices", provenance="user_task"),
            inspect_workbook(sales),
        )
        assert plan.understanding.interpretation is InterpretationVerdict.SUFFICIENTLY_CLEAR
        dedupe = next(op for op in plan.operations if op.operation.value == "remove_duplicates")
        assert dedupe.keys == ["InvoiceId"]

    def test_a_read_only_request_may_default_the_sheet(self, sales: Path) -> None:
        """Reading the largest sheet is a safe default; only mutation is gated."""
        from app.workbook import inspect_workbook

        plan = DeterministicPlanner().plan(
            UntrustedText("show me the data", provenance="user_task"), inspect_workbook(sales)
        )
        assert plan.understanding.interpretation is InterpretationVerdict.SUFFICIENTLY_CLEAR


class TestPlainEnglishColumnReferences:
    """Real requests name columns without saying "column"."""

    @pytest.mark.parametrize(
        ("phrase", "expected"),
        [
            ("trim whitespace from the notes", "Notes"),
            ("clean the customer names", "Customer"),
            ("normalise the region", "Region"),
        ],
    )
    def test_article_phrases_resolve(self, sales: Path, phrase: str, expected: str) -> None:
        from app.workbook import inspect_workbook

        plan = DeterministicPlanner().plan(
            UntrustedText(phrase, provenance="user_task"), inspect_workbook(sales)
        )
        normalize = next(
            (op for op in plan.operations if op.operation.value == "normalize_values"),
            None,
        )
        assert normalize is not None, phrase
        assert expected in normalize.columns, f"{phrase!r} -> {normalize.columns}"


class TestVerificationDoesNotMutateTheWorkbook:
    """Verification must not change the workbook it verifies.

    Indexing a worksheet by coordinate **creates** the cell in openpyxl. The
    formula-preservation check did exactly that, so asking for a coordinate in a
    deleted row resurrected that row in the in-memory model — and every data check
    that ran afterwards read the deleted rows as still present. This asserts the
    invariant directly.
    """

    def test_formula_checks_do_not_grow_the_workbook(self, tmp_path: Path) -> None:
        path = build("monthly_sales", tmp_path / "s.xlsx", rows=30)
        workbook = openpyxl.load_workbook(path)
        workbook["Sales"].delete_rows(2, 25)
        reduced = tmp_path / "reduced.xlsx"
        save_atomic(workbook, reduced)
        workbook.close()

        with opened(reduced) as loaded:
            before = loaded["Sales"].max_row
            formulas = {
                (sheet, coordinate): formula
                for sheet, coordinate, formula in formula_checks.collect_formulas(loaded)
            }
            formula_checks.static_formula_checks(loaded, before_formulas=formulas)
            after = loaded["Sales"].max_row

        assert before == after, (
            f"verification changed the sheet's row count from {before} to {after}"
        )

    def test_build_cell_index_does_not_create_cells(self, tmp_path: Path) -> None:
        path = build("monthly_sales", tmp_path / "s.xlsx", rows=10)
        with opened(path) as loaded:
            rows_before = loaded["Sales"].max_row
            index = formula_checks.build_cell_index(loaded)
            rows_after = loaded["Sales"].max_row
        assert rows_before == rows_after
        assert ("Sales", "A1") in index


class TestIntentionalRowRemovalIsNotFormulaLoss:
    """Deleting a row that carried a formula is not the same as destroying a formula.

    The formula-preservation check reported any formula that disappeared, so a
    dedupe of two duplicate rows that happened to contain a formula failed the run
    for doing exactly what it was asked to do. The executor now reports which
    coordinates it deliberately removed, and the verifier distinguishes the two.
    """

    def test_explained_removal_does_not_fail(self, tmp_path: Path) -> None:
        path = build("monthly_sales", tmp_path / "s.xlsx", rows=10)
        with opened(path) as loaded:
            before = {
                (sheet, coordinate): formula
                for sheet, coordinate, formula in formula_checks.collect_formulas(loaded)
            }
        # Remove a formula-bearing row and tell the check it was intentional.
        workbook = openpyxl.load_workbook(path)
        workbook["Sales"].delete_rows(2, 1)
        reduced = tmp_path / "reduced.xlsx"
        save_atomic(workbook, reduced)
        workbook.close()

        with opened(reduced) as loaded:
            result, lost = formula_checks.detect_formula_loss(
                before, loaded, explained_removals={("Sales", 2)}
            )
        assert lost == [], "the removed row's formula is expected, not lost"
        # Formulas below the removed row shift up, which cannot be resolved
        # without recalculation, so the check warns rather than failing.
        assert result.status.value == "warning"
        assert "moved rows" in result.message
        assert result.details["explained_by_row_removal"] == 1

    def test_no_row_removal_means_no_shift_tolerance(self, tmp_path: Path) -> None:
        """Without a row removal, a moved formula is a genuine problem."""
        path = build("monthly_sales", tmp_path / "s.xlsx", rows=10)
        with opened(path) as loaded:
            before = {
                (sheet, coordinate): formula
                for sheet, coordinate, formula in formula_checks.collect_formulas(loaded)
            }
        workbook = openpyxl.load_workbook(path)
        # Blank a row rather than deleting it, so nothing shifts.
        for column in range(1, 11):
            workbook["Sales"].cell(row=5, column=column).value = None
        edited = tmp_path / "edited.xlsx"
        save_atomic(workbook, edited)
        workbook.close()

        with opened(edited) as loaded:
            result, lost = formula_checks.detect_formula_loss(before, loaded)
        assert lost
        assert result.status.value == "failed"

    def test_unexplained_removal_still_fails(self, tmp_path: Path) -> None:
        path = build("monthly_sales", tmp_path / "s.xlsx", rows=10)
        with opened(path) as loaded:
            before = {
                (sheet, coordinate): formula
                for sheet, coordinate, formula in formula_checks.collect_formulas(loaded)
            }
        workbook = openpyxl.load_workbook(path)
        workbook["Sales"].delete_rows(2, 1)
        reduced = tmp_path / "reduced.xlsx"
        save_atomic(workbook, reduced)
        workbook.close()

        with opened(reduced) as loaded:
            result, lost = formula_checks.detect_formula_loss(before, loaded)
        assert lost, "an unexplained formula loss must still be reported"
        assert result.status.value == "failed"


class TestPreviewCountsCreatedSheetWrites:
    """The dry-run preview must not under-report.

    ``changed`` only counted cells present in both snapshots, so a plan that wrote
    a new summary sheet reported far fewer changes than it would make — and the
    planned-vs-actual anomaly check then fired on a perfectly correct run.
    """

    def test_preview_matches_actual_for_a_plan_that_creates_a_sheet(
        self, sales: Path, orchestrator: RunOrchestrator
    ) -> None:
        from app.executor import preview_plan
        from app.workbook import inspect_workbook

        inspection = inspect_workbook(sales)
        plan, *_ = orchestrator.plan(sales, UntrustedText(COMPOUND, provenance="user_task"))
        preview = preview_plan(
            list(plan.operations),
            inspection,
            run_id="run-preview",
            output_path=str(sales),
        )
        result = orchestrator.run(
            sales, UntrustedText(COMPOUND, provenance="user_task"), approve=True
        )
        actual = result.execution.total_cells_written if result.execution else 0

        # The preview is a simulation, so it should be in the right ballpark:
        # within 25% rather than off by a factor of two.
        assert preview.cells_to_change > 0
        assert abs(preview.cells_to_change - actual) <= max(2, actual * 0.25), (
            f"preview said {preview.cells_to_change}, actual was {actual}"
        )
