"""Workbook reader: the single controlled entry point for opening a workbook.

Everything that touches openpyxl goes through here. Centralising it means the
security limits, the ``keep_vba`` rule, and the read-only guarantee are enforced
in exactly one place rather than at every call site.

**This module never opens a workbook for writing.** Mutation happens in
``app.executor``, which calls :func:`save_atomic` after writing to a snapshot
(ADR-0010).
"""

from __future__ import annotations

import contextlib
import os
import warnings
from collections.abc import Iterator
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.workbook import Workbook

from app.contracts.config import LimitsConfig
from app.contracts.errors import ExcelPilotError, WorkbookSecurityError
from app.workbook.limits import (
    UnsupportedFormat,
    check_archive_integrity,
    check_extension,
    check_file_size,
)

#: Warning categories captured and reported rather than allowed to spam a CLI:
#: openpyxl emits these for some legitimate real-world files, such as a defined
#: name referring to a sheet that no longer exists.
_NOISE: tuple[type[Warning], ...] = (UserWarning, DeprecationWarning)


def load_workbook(
    path: Path,
    *,
    limits: LimitsConfig | None = None,
    read_only: bool = False,
    data_only: bool = False,
) -> Workbook:
    """Open a workbook after validating extension, size, and archive integrity.

    ``data_only=True`` returns whatever value Excel last cached rather than the
    formula text. ExcelPilot uses this **only** to detect that a cached value is
    absent or stale, and never as evidence that a formula computes correctly
    (ADR-0011).
    """
    path = Path(path)
    if not path.exists():
        raise WorkbookSecurityError(f"workbook not found: {path}", details={"path": str(path)})
    if not path.is_file():
        raise WorkbookSecurityError(f"not a file: {path}", details={"path": str(path)})

    extension = check_extension(path)
    effective = limits or LimitsConfig()
    check_file_size(path, effective)
    check_archive_integrity(path, effective)

    # XLSM must be opened with keep_vba or the macro project is silently dropped
    # on save. Enforced here so no call site can forget (ADR-0001).
    keep_vba = extension == ".xlsm"

    with warnings.catch_warnings():
        warnings.simplefilter("ignore", category=DeprecationWarning)
        try:
            return openpyxl.load_workbook(
                path,
                data_only=data_only,
                read_only=read_only,
                keep_vba=keep_vba,
                rich_text=False,
            )
        except UnsupportedFormat:
            raise
        except Exception as error:  # noqa: BLE001 - normalised into a typed error
            raise WorkbookSecurityError(
                f"could not open {path.name}: {type(error).__name__}: {error}",
                details={"path": str(path), "extension": extension},
            ) from error


@contextlib.contextmanager
def opened(
    path: Path,
    *,
    limits: LimitsConfig | None = None,
    data_only: bool = False,
) -> Iterator[Workbook]:
    """Open a workbook and always close it.

    Used by inspection, diff, and verification — the read paths. The executor
    deliberately does not use this, because it must hold a workbook open across
    a save.
    """
    workbook = load_workbook(path, limits=limits, data_only=data_only)
    try:
        yield workbook
    finally:
        with contextlib.suppress(Exception):
            workbook.close()


def capture_warnings(path: Path, *, limits: LimitsConfig | None = None) -> list[str]:
    """Open and close a workbook, returning any warnings it produced.

    Real workbooks routinely carry a defined name pointing at a deleted sheet or
    an unsupported extension. That is worth telling an operator about rather than
    hiding, so the caller can surface it.

    Best-effort by design: if the workbook cannot be opened at all this returns
    an empty list rather than raising, because it is called for supplementary
    context. The real open, in :func:`load_workbook`, does raise.
    """
    messages: list[str] = []
    with warnings.catch_warnings(record=True) as captured:
        warnings.simplefilter("always")
        try:
            workbook = load_workbook(path, limits=limits)
            workbook.close()
        except ExcelPilotError:
            return messages
    messages.extend(
        f"{warning.category.__name__}: {warning.message}"
        for warning in captured
        if issubclass(warning.category, _NOISE)
    )
    return messages


def save_atomic(workbook: Workbook, destination: Path, *, fsync: bool = True) -> Path:
    """Save a workbook so no reader ever sees a partial file.

    Writes to a temporary file in the *destination* directory (so the final move
    is within one filesystem and therefore atomic), optionally fsyncs, then
    ``os.replace``s into position. A crash mid-write cannot corrupt an existing
    file, and a reader never observes a half-written workbook (ADR-0010).
    """
    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    # Named with the pid so two concurrent processes cannot collide, and so an
    # interrupted run leaves an obviously-temporary file behind.
    temporary = destination.with_name(f".{destination.name}.tmp-{os.getpid()}")
    try:
        workbook.save(temporary)
        if fsync:
            _fsync_file(temporary)
        # Path.replace is os.replace semantics: atomic within a filesystem.
        temporary.replace(destination)
    except Exception:
        with contextlib.suppress(OSError):
            temporary.unlink()
        raise
    return destination


def _fsync_file(path: Path) -> None:
    """Flush a file's contents to disk before the atomic rename."""
    try:
        handle = os.open(path, os.O_RDONLY)
    except OSError:
        return
    try:
        os.fsync(handle)
    except OSError:
        pass
    finally:
        os.close(handle)


def copy_snapshot(source: Path, destination: Path) -> Path:
    """Byte-for-byte copy of a workbook, used to snapshot the source before work.

    A copy, not a re-save: re-serialising would change the bytes and the
    shared-string cache (ADR-0009), so a "snapshot" produced by re-saving would
    not be the user's original.
    """
    import shutil

    source, destination = Path(source), Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, destination)
    return destination


def sheet_state(sheet: Any) -> str:
    """Worksheet visibility as a plain string, defaulting safely."""
    return str(getattr(sheet, "sheet_state", "visible") or "visible")


def has_vba(path: Path) -> bool:
    """Whether the package actually contains a VBA project.

    Checking the zip rather than trusting the extension: a ``.xlsm`` saved
    without macros is common, and the fixture set contains several. This is what
    lets the inspector report ``has_vba`` truthfully.
    """
    import zipfile

    try:
        with zipfile.ZipFile(path) as archive:
            return any("vbaproject.bin" in name.lower() for name in archive.namelist())
    except (OSError, zipfile.BadZipFile):
        return False


def has_external_links(path: Path) -> bool:
    """Whether the workbook links to other workbooks.

    Such links can trigger refresh prompts and reach outside the workspace, so
    policy and the inspector both treat their presence as worth reporting.
    """
    import zipfile

    try:
        with zipfile.ZipFile(path) as archive:
            return any(
                name.startswith("xl/externalLinks/") and name.endswith(".xml")
                for name in archive.namelist()
            )
    except (OSError, zipfile.BadZipFile):
        return False


__all__ = [
    "capture_warnings",
    "copy_snapshot",
    "has_external_links",
    "has_vba",
    "load_workbook",
    "opened",
    "save_atomic",
    "sheet_state",
]
