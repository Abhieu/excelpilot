"""Anomaly detection — deterministic, and attributed.

Every anomaly records its ``source`` (``deterministic`` | ``model`` | ``jev`` |
``human``), so a probabilistic finding is never presented as a verified fact
(spec section 21).

The scope is deliberately bounded. These are the checks that deterministic
computation answers well. A model or JEV may add interpretation, but only as a
clearly-labelled advisory finding, never as a measurement.
"""

from __future__ import annotations

import statistics
from typing import Any

from openpyxl.workbook import Workbook

from app.contracts.config import AnomalyConfig
from app.contracts.enums import AnomalyKind, AnomalySeverity, AnomalySource
from app.contracts.operations import Target
from app.contracts.verification import Anomaly
from app.contracts.verification import CheckResult as _CheckResult  # noqa: F401 - re-export guard
from app.workbook.table import read_table


def detect(
    before_workbook: Workbook | None,
    after_workbook: Workbook,
    *,
    planned_cells_changed: int = 0,
    actual_cells_changed: int = 0,
    config: AnomalyConfig | None = None,
) -> list[Anomaly]:
    """Run every deterministic anomaly check."""
    settings = config or AnomalyConfig()
    found: list[Anomaly] = []
    found.extend(_structural_anomalies(before_workbook, after_workbook))
    found.extend(_row_count_anomalies(before_workbook, after_workbook, settings))
    found.extend(_divergence_anomaly(planned_cells_changed, actual_cells_changed, settings))
    if before_workbook is not None:
        found.extend(_duplicate_rate_anomalies(before_workbook, after_workbook, settings))
        found.extend(_null_rate_anomalies(before_workbook, after_workbook, settings))
        found.extend(_outliers(after_workbook, settings))
    return found


def _structural_anomalies(
    before_workbook: Workbook | None, after_workbook: Workbook
) -> list[Anomaly]:
    """Sheet-level structural surprises."""
    if before_workbook is None:
        return []
    removed = [n for n in before_workbook.sheetnames if n not in after_workbook.sheetnames]
    added = [n for n in after_workbook.sheetnames if n not in before_workbook.sheetnames]
    found: list[Anomaly] = []
    if removed:
        found.append(
            Anomaly(
                kind=AnomalyKind.STRUCTURAL_CHANGE,
                severity=AnomalySeverity.ERROR,
                source=AnomalySource.DETERMINISTIC,
                message=f"{len(removed)} sheet(s) were removed: {', '.join(removed)}",
                evidence={"removed": removed},
                detector="sheet_presence",
            )
        )
    if added:
        found.append(
            Anomaly(
                kind=AnomalyKind.STRUCTURAL_CHANGE,
                severity=AnomalySeverity.INFO,
                source=AnomalySource.DETERMINISTIC,
                message=f"{len(added)} sheet(s) were added: {', '.join(added)}",
                evidence={"added": added},
                detector="sheet_presence",
            )
        )
    return found


def _row_count_anomalies(
    before_workbook: Workbook | None,
    after_workbook: Workbook,
    config: AnomalyConfig,
) -> list[Anomaly]:
    """Row counts that moved by more than expected."""
    if before_workbook is None:
        return []
    found: list[Anomaly] = []
    for title in before_workbook.sheetnames:
        if title not in after_workbook.sheetnames:
            continue
        before_rows = before_workbook[title].max_row or 0
        after_rows = after_workbook[title].max_row or 0
        if before_rows <= 0:
            continue
        change = abs(after_rows - before_rows) / before_rows
        if change > config.row_count_change_ratio:
            direction = "decreased" if after_rows < before_rows else "increased"
            found.append(
                Anomaly(
                    kind=AnomalyKind.ROW_COUNT_CHANGE,
                    severity=(
                        AnomalySeverity.ERROR
                        if after_rows < before_rows and change > 0.5
                        else AnomalySeverity.WARNING
                    ),
                    source=AnomalySource.DETERMINISTIC,
                    message=(
                        f"row count on {title!r} {direction} from {before_rows:,} to "
                        f"{after_rows:,} ({change:.0%})"
                    ),
                    sheet=title,
                    evidence={
                        "before": before_rows,
                        "after": after_rows,
                        "change_ratio": round(change, 4),
                    },
                    detector="row_count_delta",
                )
            )
    return found


def _divergence_anomaly(planned: int, actual: int, config: AnomalyConfig) -> list[Anomaly]:
    """Planned-versus-actual divergence.

    This is how an operation behaving unexpectedly gets caught. A large gap
    between what the plan said it would change and what actually changed is
    itself the finding, and it is reported even when the run "succeeded".
    """
    if planned <= 0 or actual <= 0:
        return []
    divergence = abs(actual - planned) / planned
    if divergence <= config.planned_vs_actual_divergence:
        return []
    return [
        Anomaly(
            kind=AnomalyKind.PLANNED_VS_ACTUAL_DIVERGENCE,
            severity=AnomalySeverity.WARNING,
            source=AnomalySource.DETERMINISTIC,
            message=(
                f"the plan predicted {planned:,} changed cells but {actual:,} changed "
                f"({divergence:.0%} divergence)"
            ),
            evidence={"planned": planned, "actual": actual, "divergence": round(divergence, 4)},
            detector="planned_vs_actual",
        )
    ]


def _duplicate_rate_anomalies(
    before_workbook: Workbook,
    after_workbook: Workbook,
    config: AnomalyConfig,
) -> list[Anomaly]:
    """A spike in duplicates suggests a bad join or a bad dedupe key."""
    found: list[Anomaly] = []
    for title in after_workbook.sheetnames:
        if title not in before_workbook.sheetnames:
            continue
        before_rate = _max_duplicate_rate(before_workbook, title)
        after_rate = _max_duplicate_rate(after_workbook, title)
        if before_rate is None or after_rate is None:
            continue
        if after_rate - before_rate > config.duplicate_rate_increase:
            found.append(
                Anomaly(
                    kind=AnomalyKind.DUPLICATE_SPIKE,
                    severity=AnomalySeverity.WARNING,
                    source=AnomalySource.DETERMINISTIC,
                    message=(
                        f"duplicate rate on {title!r} rose from {before_rate:.1%} to "
                        f"{after_rate:.1%}"
                    ),
                    sheet=title,
                    evidence={"before": before_rate, "after": after_rate},
                    detector="duplicate_rate_delta",
                )
            )
    return found


def _null_rate_anomalies(
    before_workbook: Workbook,
    after_workbook: Workbook,
    config: AnomalyConfig,
) -> list[Anomaly]:
    """Blank cells appearing where there were values."""
    found: list[Anomaly] = []
    for title in after_workbook.sheetnames:
        if title not in before_workbook.sheetnames:
            continue
        before_rate = _overall_null_rate(before_workbook, title)
        after_rate = _overall_null_rate(after_workbook, title)
        if before_rate is None or after_rate is None:
            continue
        if after_rate - before_rate > config.null_rate_increase:
            found.append(
                Anomaly(
                    kind=AnomalyKind.NULL_RATE_INCREASE,
                    severity=AnomalySeverity.WARNING,
                    source=AnomalySource.DETERMINISTIC,
                    message=(
                        f"null rate on {title!r} rose from {before_rate:.1%} to {after_rate:.1%}"
                    ),
                    sheet=title,
                    evidence={"before": before_rate, "after": after_rate},
                    detector="null_rate_delta",
                )
            )
    return found


def _outliers(workbook: Workbook, config: AnomalyConfig) -> list[Anomaly]:
    """Extreme values in a numeric column, by z-score.

    Bounded to one finding per column naming the worst offenders, rather than one
    per cell, so a column with a legitimate long tail does not flood the report.
    """
    found: list[Anomaly] = []
    for title in workbook.sheetnames:
        try:
            view = read_table(workbook, Target(sheet=title))
        except Exception:  # noqa: BLE001 - an unreadable sheet is not this check's business
            continue
        for header in view.headers:
            if not header:
                continue
            try:
                values = [
                    float(value)
                    for value in view.column_values(header)
                    if isinstance(value, (int, float)) and not isinstance(value, bool)
                ]
            except Exception:  # noqa: BLE001 - the column is not resolvable
                continue
            if len(values) < 8:
                continue
            mean = statistics.fmean(values)
            try:
                deviation = statistics.pstdev(values)
            except statistics.StatisticsError:
                continue
            if deviation == 0:
                continue
            outliers = [
                (index, value)
                for index, value in enumerate(values)
                if abs(value - mean) / deviation > config.outlier_z_score
            ]
            if outliers:
                worst = sorted(outliers, key=lambda pair: abs(pair[1] - mean), reverse=True)[:3]
                found.append(
                    Anomaly(
                        kind=AnomalyKind.OUTLIER_VALUE,
                        severity=AnomalySeverity.INFO,
                        source=AnomalySource.DETERMINISTIC,
                        message=(
                            f"{len(outliers)} outlier value(s) in {title}.{header} "
                            f"(mean {mean:,.2f}); largest: "
                            + ", ".join(f"{value:,.2f}" for _, value in worst)
                        ),
                        sheet=title,
                        evidence={
                            "column": header,
                            "mean": mean,
                            "stdev": deviation,
                            "count": len(outliers),
                        },
                        detector="z_score_outlier",
                    )
                )
    return found


def _max_duplicate_rate(workbook: Workbook, title: str) -> float | None:
    """Highest duplicate rate across the sheet's columns."""
    try:
        view = read_table(workbook, Target(sheet=title))
    except Exception:  # noqa: BLE001
        return None
    if not view.headers or not view.rows:
        return None
    rates: list[float] = []
    for header in view.headers:
        if not header:
            continue
        try:
            values = view.column_values(header)
        except Exception:  # noqa: BLE001
            continue
        if not values:
            continue
        seen: set[str] = set()
        duplicates = 0
        for value in values:
            key = "" if value is None else str(value).strip().lower()
            if key in seen:
                duplicates += 1
            else:
                seen.add(key)
        rates.append(duplicates / len(values))
    return max(rates) if rates else None


def _overall_null_rate(workbook: Workbook, title: str) -> float | None:
    """Fraction of blank cells across the sheet's data area."""
    worksheet = workbook[title]
    total = 0
    blanks = 0
    for row in worksheet.iter_rows():
        for cell in row:
            if cell.row == 1:
                continue  # header
            total += 1
            if cell.value is None or (isinstance(cell.value, str) and not cell.value.strip()):
                blanks += 1
    return blanks / total if total else None


def from_model(text: str, *, evidence: dict[str, Any] | None = None) -> Anomaly:
    """Wrap a model-produced observation as an explicitly attributed anomaly.

    Exists so a model finding can never be recorded with
    ``source=deterministic`` by accident. Severity is capped at ``warning``
    because a probabilistic observation is not a measurement.
    """
    return Anomaly(
        kind=AnomalyKind.UNCLASSIFIED,
        severity=AnomalySeverity.WARNING,
        source=AnomalySource.MODEL,
        message=text,
        evidence=evidence or {},
        detector="model_observation",
    )


def from_jev(text: str, *, evidence: dict[str, Any] | None = None) -> Anomaly:
    """Wrap a JEV-produced observation as an explicitly attributed anomaly."""
    return Anomaly(
        kind=AnomalyKind.UNCLASSIFIED,
        severity=AnomalySeverity.WARNING,
        source=AnomalySource.JEV,
        message=text,
        evidence=evidence or {},
        detector="jev_observation",
    )


__all__ = ["detect", "from_jev", "from_model"]
